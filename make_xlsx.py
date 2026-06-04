# -*- coding: utf-8 -*-
"""待查.xlsx — 四人分派，底纹颜色区分"""
import csv, os
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

os.chdir(r"D:\Desktop\VAST CHALLENGE")

with open("待查.csv", encoding="utf-8-sig") as f:
    rows = list(csv.reader(f))
header = rows[0]
data = rows[1:]

# Four person colors
COLORS = {
    "甲": PatternFill(start_color="FFD6E4F0", end_color="FFD6E4F0", fill_type="solid"),  # light blue
    "乙": PatternFill(start_color="FFE2EFDA", end_color="FFE2EFDA", fill_type="solid"),  # light green
    "丙": PatternFill(start_color="FFFFF2CC", end_color="FFFFF2CC", fill_type="solid"),  # light yellow
    "丁": PatternFill(start_color="FFFCE4D6", end_color="FFFCE4D6", fill_type="solid"),  # light orange
}
HEADER_FILL = PatternFill(start_color="FF4472C4", end_color="FF4472C4", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
thin_border = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin'))

# Sort: by creation_type, then by year within each type
type_order = {"作画": 0, "书法": 1, "篆刻": 2}
data.sort(key=lambda r: (type_order.get(r[1], 99), r[3]))

# Distribute: interleave by 4 within each type group, so everyone gets a balanced mix
from collections import defaultdict
by_type = defaultdict(list)
for r in data:
    by_type[r[1]].append(r)

assigned = []
for ctype in ["作画", "书法", "篆刻"]:
    group = by_type[ctype]
    for i, r in enumerate(group):
        person = ["甲", "乙", "丙", "丁"][i % 4]
        assigned.append((person, r))

# Write xlsx
wb = Workbook()
ws = wb.active
ws.title = "待查作品"

# Header
for ci, h in enumerate(header + ["负责人"], 1):
    cell = ws.cell(row=1, column=ci, value=h)
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = Alignment(horizontal='center')
    cell.border = thin_border

# Data rows
for ri, (person, r) in enumerate(assigned, 2):
    fill = COLORS[person]
    for ci, val in enumerate(r + [person], 1):
        cell = ws.cell(row=ri, column=ci, value=val)
        if ci == 1:  # 作品名列 — colored
            cell.fill = fill
        cell.border = thin_border
        cell.alignment = Alignment(vertical='center')

# Column widths
ws.column_dimensions['A'].width = 42  # 作品名
ws.column_dimensions['B'].width = 10  # 创作类型
ws.column_dimensions['C'].width = 8   # 题材
ws.column_dimensions['D'].width = 8   # 年份
ws.column_dimensions['E'].width = 35  # 别名
ws.column_dimensions['F'].width = 14  # event_id
ws.column_dimensions['G'].width = 60  # 原文
ws.column_dimensions['H'].width = 8   # 负责人

# Freeze header row
ws.freeze_panes = 'A2'

# Auto-filter
ws.auto_filter.ref = f"A1:H{len(assigned)+1}"

wb.save("待查.xlsx")

# Stats
from collections import Counter
person_counts = Counter(p for p, _ in assigned)
print(f"待查.xlsx: {len(assigned)} rows")
for p in ["甲","乙","丙","丁"]:
    types = Counter(r[1] for per, r in assigned if per == p)
    print(f"  {p}: {person_counts[p]}件 — {dict(types)}")
